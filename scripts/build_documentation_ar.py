#!/usr/bin/env python3
"""Build the professional Arabic documentation deliverable (DOCX, RTL).

Covers the three course requirements verbatim, with real measured values from
the verified 36-step evidence run and real screenshots.
Usage: python scripts/build_documentation_ar.py
Writes: /home/user/StegoNexus_Documentation_AR.docx
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
EVID = ROOT / "artifacts" / "evidence_report"
SHOTS = ROOT / "presentation" / "screenshots"
TMP = ROOT / "data" / "doc_imgs"
OUT = "/home/user/StegoNexus_Documentation_AR.docx"

DARK = RGBColor(0x0D, 0x3B, 0x66)
ACCENT = RGBColor(0x0B, 0x7A, 0x3D)
CYAN = RGBColor(0x0E, 0x74, 0x90)
RED = RGBColor(0xB3, 0x26, 0x1E)
GRAY = RGBColor(0x44, 0x55, 0x66)
AMBER = RGBColor(0x8A, 0x6D, 0x00)

STEPS = json.loads((EVID / "evidence_steps.json").read_text(encoding="utf-8"))
BY_FILE = {s["file"]: s for s in STEPS}
TMP.mkdir(parents=True, exist_ok=True)


# ------------------------------------------------------------------ helpers
def set_rtl(p) -> None:
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)


def srun(run, size=11, bold=False, color=None, mono=False) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = "Consolas" if mono else "Segoe UI"
    rPr = run._element.get_or_add_rPr()
    rF = rPr.find(qn("w:rFonts"))
    if rF is None:
        rF = OxmlElement("w:rFonts")
        rPr.append(rF)
    rF.set(qn("w:cs"), "Consolas" if mono else "Segoe UI")
    sz = OxmlElement("w:szCs"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    if bold:
        b = OxmlElement("w:bCs"); b.set(qn("w:val"), "1"); rPr.append(b)
    r = OxmlElement("w:rtl"); r.set(qn("w:val"), "1"); rPr.append(r)


def H(doc, text, size=16, color=DARK, before=14, after=4, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph()
    set_rtl(p)
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    srun(p.add_run(text), size=size, bold=True, color=color)
    return p


def P(doc, text, size=11, bold=False, color=None, bullet=False, mono=False) -> None:
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    set_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(3)
    # allow simple **bold** segments
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        srun(p.add_run(part), size=size, bold=bold or (i % 2 == 1),
             color=color, mono=mono)


def TBL(doc, headers, rows, widths=None, size=10) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        set_rtl(p)
        srun(p.add_run(htxt), size=size, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            set_rtl(p)
            srun(p.add_run(str(val)), size=size)
    if widths:
        for i, w in enumerate(widths):
            for r in t.rows:
                r.cells[i].width = Inches(w)
    doc.add_paragraph()


def IMG(doc, filename, caption, folder=EVID, width=6.1) -> None:
    src = Path(folder) / filename
    if not src.exists():
        return
    im = Image.open(src).convert("RGB")
    new_w = int(width * 96)
    new_h = max(1, int(im.height * new_w / im.width))
    im = im.resize((new_w, new_h))
    dst = TMP / (Path(filename).stem + ".jpg")
    im.save(dst, "JPEG", quality=82)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(dst), width=Inches(width))
    cap = doc.add_paragraph()
    set_rtl(cap)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun(cap.add_run("📷 " + caption), size=9.5, bold=False, color=GRAY)


def step_rows(filename):
    s = BY_FILE[filename]
    return s["inputs"], s["outputs"], s["explain"]


def REAL(doc, filename, note="القياسات الفعلية لهذه الخطوة (من التشغيل الموثق)"):
    inp, outp, expl = step_rows(filename)
    TBL(doc, ["المدخلات الحقيقية", "القيمة"], inp, widths=[2.6, 4.2])
    TBL(doc, ["النواتج المقاسة", "القيمة"], outp, widths=[2.6, 4.2])
    P(doc, expl, size=10, color=GRAY)


# ------------------------------------------------------------------ document
doc = Document()
for sec in doc.sections:
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

# ============ الغلاف
for _ in range(3):
    doc.add_paragraph()
H(doc, "مشروع مادة إخفاء البيانات والتحليل الجنائي الرقمي", size=14, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
H(doc, "StegoNexus", size=34, align=WD_ALIGN_PARAGRAPH.CENTER)
H(doc, "منصة متكاملة لإخفاء واستخراج البيانات وتحليل البرمجيات الخبيثة والتحليل الجنائي الرقمي",
  size=15, color=CYAN, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
H(doc, "وثيقة التوثيق الشاملة", size=20, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
doc.add_paragraph()
H(doc, "تطوير وإعداد: Mohammed Moneer Al-absi", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
H(doc, "إشراف: أسامة الشلالي", size=13, color=CYAN, align=WD_ALIGN_PARAGRAPH.CENTER)
H(doc, "العام الجامعي 2026", size=12, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ============ المحتويات
H(doc, "فهرس المحتويات", size=17)
toc = [
    "1. مقدمة: المشكلة والحل والأهداف",
    "2. نظرة عامة على الأداة والبنية المعمارية",
    "3. المصطلحات الرئيسية (Key Terminology)",
    "4. المتطلب الأول: أداة عرض وحقن البيانات الوصفية",
    "5. المتطلب الثاني: أدوات الإخفاء والاستخراج في القضايا (1، 2، 3)",
    "6. المتطلب الثالث: أدوات الإخفاء والاستخراج في الوسائط",
    "      6.1 الصور — تقنيتان   6.2 الصوت — ثلاث تقنيات   6.3 الفيديو — تقنيتان",
    "      6.4 الشبكة — قناة IPv4-ID   6.5 البرامج الضارة والفيروسات — تقنيتان + الكشف",
    "7. الأمان والأخلاقيات",
    "8. الاختبارات والتحقق من الصحة",
    "9. خريطة المتطلبات ← التنفيذ ← التحقق",
    "10. التثبيت والتشغيل وإعادة إنتاج النتائج",
    "11. الخاتمة",
]
for t in toc:
    P(doc, t, size=11.5)
doc.add_page_break()

# ============ 1 مقدمة
H(doc, "1. مقدمة", size=17)
H(doc, "1.1 المشكلة", size=13)
P(doc, "أدوات إخفاء البيانات المتاحة للمقرر متفرقة ومحدودة: كل أداة تنفذ تقنية واحدة غالبًا على منصة واحدة "
       "(OpenPuff وDeepSound وCyberHide تعمل على Windows حصرًا)، ولا يوجد ربط بين عمليات الإخفاء والاستخراج "
       "وبين التحليل الجنائي وإدارة القضايا وتوليد التقارير، فضلًا عن صعوبة إثبات أن نتائج هذه الأدوات حقيقية.", size=11)
H(doc, "1.2 الحل: StegoNexus", size=13)
P(doc, "منصة موحدة تعمل على **Kali Linux** مبنية بلغة **Python** وواجهة **PySide6**، تنفذ كل تقنيات الإخفاء "
       "المدروسة في المقرر بمحركات أصلية مكتوبة خصيصًا، وتدمج الأدوات الخارجية الحقيقية (ExifTool، Steghide، "
       "Binwalk، Foremost، FFmpeg، tshark، zsteg)، وتضيف طبقة إدارة تحقيق كاملة: قضايا وأدلة بتجزئاتها "
       "وسجل تدقيق وتقارير بثلاث صيغ.", size=11)
H(doc, "1.3 الأهداف", size=13)
for t in ["تغطية 100% من متطلبات المقرر: عرض وحقن البيانات الوصفية، والإخفاء والاستخراج في النص والصورة "
          "والصوت والفيديو والشبكة والملفات التنفيذية (البرامج الضارة/الفيروسات).",
          "تنفيذ كل وسيط بأكثر من تقنية (الصوت 3 تقنيات، الصورة والفيديو والتنفيذية تقنيتان لكل منها).",
          "إثبات كل نتيجة قياسًا: تطابق SHA-256 بين المسترجَع والأصل في كل جولة إخفاء/استخراج.",
          "أمانة أكاديمية كاملة: ملصق حالة لكل خاصية، وسجل تدقيق لكل إجراء، وإعلان صريح للحدود."]:
    P(doc, t, bullet=True, size=11)
H(doc, "1.4 نطاق الاستخدام", size=13)
P(doc, "أداة دفاعية وتعليمية للاستخدام المخبري المصرَّح به فقط: لا تنفيذ للعينات الخبيثة، ولا أي وظيفة هجومية، "
       "والإرسال الشبكي مقيد بعنوان الحلقة المحلية والنطاقات الخاصة.", size=11)

# ============ 2 نظرة عامة
H(doc, "2. نظرة عامة على الأداة", size=17)
H(doc, "2.1 البنية المعمارية", size=13)
TBL(doc, ["الطبقة", "المحتوى"],
    [["الواجهات", "PySide6 GUI — 20 شاشة (ثيم داكن/فاتح، عمليات غير متجمدة على خيوط مستقلة) + واجهة CLI تفاعلية؛ كلاهما فوق نفس طبقة الخدمات"],
     ["خدمات (17 وحدة)", "كل وحدة Facade: تحقق مدخلات ← تنفيذ ← تسجيل صف تحليل + إدخال سجل تدقيق ← إرجاع نتيجة غنية"],
     ["محركات الإخفاء (10)", "lsb_text · lsb_image · lsb_audio · phase_coding · spread_spectrum · container(EOF) · video_lsb · video_spread · pe_stego(overlay/slack)"],
     ["الأدوات الخارجية", "ToolExecutor آمن (بدون shell=True): exiftool · steghide · binwalk · foremost · ffmpeg/ffprobe · tshark · file · strings · zsteg"],
     ["التحليل الدفاعي", "فحص PE/ELF/Office · IOC · كتالوغ تقنيات التمويه · التحليل الساكن للبرمجيات الخبيثة"],
     ["التخزين", "SQLite (WAL): قضايا · أدلة وتجزئاتها · نتائج · تحليلات · تقارير · سجلات — لوحة المعلومات محسوبة منه كليًا"]],
    widths=[1.7, 5.1])
IMG(doc, "dashboard", "لوحة المعلومات — كل الأرقام محسوبة من قاعدة التحقيق الحية", folder=SHOTS)
H(doc, "2.2 سياسة الصدق (لغة الحالات)", size=13)
TBL(doc, ["الحالة", "المعنى", "أمثلة"],
    [["IMPLEMENTED", "محرك أصلي كتبناه واختبرناه", "10 محركات إخفاء + التحليل الجنائي"],
     ["INTEGRATED", "أداة خارجية حقيقية مدمجة", "ExifTool · Steghide · FFmpeg · Binwalk"],
     ["REFERENCE", "أداة مقرر نشرحها ولا نزيّفها", "OpenPuff · CyberHide · DeepSound · Coagula"],
     ["EXTERNAL", "أداة نظام يمكن إطلاقها", "Wireshark · Audacity"],
     ["UNAVAILABLE", "غير مثبتة — تُعلن بصدق", "أي أداة غائبة تظهر هكذا دون نتائج ملفقة"]],
    widths=[1.5, 2.7, 2.6])
H(doc, "2.3 الإطار الجنائي", size=13)
P(doc, "كل عمل يبدأ بقضية (Case) وتُستورد الأدلة (Evidence) فتُحسب لها 4 تجزئات (MD5/SHA-1/SHA-256/SHA-512) "
       "وتُفتح سلسلة العهدة، والنتائج (Findings) وسجل التدقيق (Audit Log) والتقارير كلها مرتبطة بالقضية. "
       "الدليل الأصلي **لا يُعدَّل أبدًا** — كل العمليات على نسخ عمل.", size=11)
IMG(doc, "s02.png", "إنشاء قضية حقيقية واستيراد 6 أدلة بتجزئاتها — من التشغيل الموثق")

# ============ 3 terminology
H(doc, "3. المصطلحات الرئيسية (Key Terminology)", size=17)
TBL(doc, ["Term (EN)", "الشرح"],
    [["Steganography / Steganalysis", "إخفاء البيانات داخل وسائط رقمية / علم كشف ذلك واستخراجه"],
     ["LSB — Least Significant Bit", "البت الأقل أهمية: تعديله يغير القيمة ±1 فقط (غير محسوس)"],
     ["Carrier / Cover — Payload", "الوسيط الحامل — الحمولة السرية المخفاة"],
     ["SNR — Signal-to-Noise Ratio", "نسبة الإشارة إلى الضوضاء: مقياس التشويه بعد الإخفاء (dB)"],
     ["Entropy (Shannon)", "الإنتروبيا: مقياس العشوائية 0–8 bits/byte؛ العالية مؤشر ضغط/تشفير"],
     ["Phase Coding", "ترميز الطور: إخفاء بتات في طور التحول الفوريئ FFT"],
     ["DSSS — Direct Sequence Spread Spectrum", "طيف الانتشار المباشر: نشر كل بت عبر متتالية ضوضاء زائفة PN"],
     ["Covert Channel", "قناة خفية عبر حقول بروتوكولات (مثل IPv4 Identification)"],
     ["EOF Container / PE Overlay", "حاوية ملحقة بعد نهاية الملف / بعد أقسام الملف التنفيذي"],
     ["Slack Space", "الحشو غير المستخدم بين VirtualSize وSizeOfRawData في أقسام PE"],
     ["AES-256-GCM / PBKDF2", "تشفير مصادق + اشتقاق مفتاح (200,000 تكرار في هذه الأداة)"],
     ["Hashing (MD5/SHA-1/SHA-256/SHA-512)", "التجزئة: إثبات السلامة والمطابقة"],
     ["Metadata (EXIF)", "البيانات الوصفية: معلومات عن الملف لا محتواه"],
     ["IOC — Indicators of Compromise", "مؤشرات الاختراق (روابط، عناوين، سلاسل مريبة)"],
     ["Static Analysis", "التحليل الساكن: فحص البايتات دون تنفيذ العينة"],
     ["Carving", "استخراج الملفات المضمنة بالتوقيعات (binwalk/foremost)"],
     ["Timestomping / Masquerading", "تزوير التوقيتات / انتحال نوع الملف"]],
    widths=[3.0, 3.8], size=9.5)

# ============ 4 metadata
H(doc, "4. المتطلب الأول: أداة عرض البيانات الوصفية وحقنها", size=17)
H(doc, "4.1 المفهوم", size=13)
P(doc, "البيانات الوصفية (Metadata) هي معلومات «عن» الملف لا محتواه: الكاميرا والعدسة والإحداثيات الجغرافية، "
       "البرنامج المنتج، التواريخ، التعليقات. لها قيمة جنائية مزدوجة: تثبت المنشأ والتوقيت، وقد تسرّب معلومات "
       "حساسة، كما يمكن استخدام حقنها قناةَ إخفاءٍ خفيفة.", size=11)
H(doc, "4.2 الأداة المستخدمة والتقنيات", size=13)
P(doc, "الأداة: **ExifTool** الخارجية الحقيقية (حالة INTEGRATED) عبر شاشة Metadata في الواجهة، وتدعم أربع عمليات:", size=11)
TBL(doc, ["العملية", "ماذا تفعل", "ضمانة الأمان"],
    [["Read", "عرض كل وسوم الملف (المجموعة/الوسم/القيمة)", "قراءة فقط"],
     ["Inject", "كتابة وسوم (Comment/Artist/Copyright/…) في نسخة جديدة", "الأصل لا يُمَس — يثبت قبل/بعد SHA-256"],
     ["Compare", "فرق قبل/بعد وسمًا بوسم: مضاف/محذوف/معدَّل", "—"],
     ["Strip", "نسخة نظيفة بلا وسوم حساسة للمشاركة الآمنة", "الإزالة على النسخة"]],
    widths=[1.2, 3.6, 2.0])
H(doc, "4.3 النتائج الحقيقية من التشغيل الموثق", size=13)
IMG(doc, "s05.png", "قراءة وسوم JPEG — 32 وسمًا حقيقيًا عبر exiftool")
REAL(doc, "s05.png")
IMG(doc, "s06.png", "حقن 3 وسوم في نسخة — سطر قبل/بعد يثبت أن الأصل لم يتغير")
REAL(doc, "s06.png")
IMG(doc, "s07.png", "إزالة البيانات الوصفية (Strip) — نسخة نظيفة للمشاركة")
REAL(doc, "s07.png")
P(doc, "الاستنتاج الموثق: الحقن تم في نسخة فقط (تجزئة الأصل متطابقة قبل/بعد) — وهذا هو الضمان الأمني المطلوب في المناقشة.",
  size=11, bold=True, color=ACCENT)

# ============ 5 cases
H(doc, "5. المتطلب الثاني: أدوات الإخفاء والاستخراج في القضايا (1، 2، 3)", size=17)
P(doc, "تدير الأداة القضايا عبر شاشتي Cases/Evidence: إنشاء قضية ← استيراد أدلة بتجزئاتها ← تنفيذ عمليات الإخفاء/الاستخراج "
       "والتحليل ← تسجيل النتائج ← توليد تقرير PDF/HTML/JSON. وتضم الأداة أربع قضايا تدريبية جاهزة موسومة "
       "TRAINING-SYNTHETIC تعكس أنماط القضايا الثلاث في المقرر وتثبت عمل كل الأدوات داخل سير قضية كامل:", size=11)
TBL(doc, ["القضية", "النمط في المقرر", "الأدوات والتقنيات المستخدمة داخلها", "التحقق المقاس"],
    [["القضية الأولى (نمط A)", "صورة + ميتاداتا",
      "Metadata Inject/Compare (ExifTool) + Image LSB الأصلي (+ Steghide متاحة)",
      "3 وسوم مضافة · استرجاع مطابق SHA-256 · الأصل سليم"],
     ["القضية الثانية (نمط B)", "صوت",
      "Audio LSB (وأيضًا Phase Coding وDSSS متاحتان في نفس الشاشة)",
      "SNR ~97dB · استرجاع مطابق"],
     ["القضية الثالثة (نمط C)", "شبكة + برنامج خبيث",
      "Network PCAP Interpreter + Malware Static Analysis (ساكن دفاعي)",
      "ثقة High للاعتراض (مؤشر لا دليل) · تحليل دون تنفيذ"],
     ["القضية الداعمة (نمط D)", "ملف تنفيذي (برامج ضارة)",
      "Executable Stego: Overlay + Slack + Scan + Static flag",
      "الحجم لم يتغير (slack) · استرجاع مطابق · إلواح تلقائي"]],
    widths=[1.35, 1.35, 2.85, 1.55], size=9.5)
P(doc, "قضايا المقرر الأصلية: تُستورد كما هي عبر Cases ← New Case ← Evidence ← Import ثم تُنفذ بنفس الشاشات — "
       "الخدمات مستقلة عن صيغة الملف ولا يحتاج ذلك أي تعديل في الكود.", size=11, color=AMBER)
IMG(doc, "s35.png", "تقرير القضية بثلاث صيغ (PDF/HTML/JSON) مولد من سجلات القضية نفسها")

# ============ 6 media
H(doc, "6. المتطلب الثالث: أدوات الإخفاء والاستخراج في الوسائط", size=17)
P(doc, "لكل وسيط اخترنا أكثر من تقنية كما يسمح المتطلب، ونعرض لكل تقنية: المبدأ، خطوات الخوارزمية كما نُفذت فعلًا "
       "في الكود، ثم القياسات الحقيقية من التشغيل الموثق (كل جولة = إخفاء ← استخراج ← مقارنة SHA-256).", size=11)

# ---- 6.1 images
H(doc, "6.1 الإخفاء في الصور — تقنيتان", size=14, color=CYAN)
H(doc, "التقنية 1: LSB الأصلي (محرك IMPLEMENTED)", size=12)
P(doc, "المبدأ: استبدال البت الأقل أهمية في قنوات البكسل — يجب أن يكون الحامل والمخرج بلا فقد (PNG/BMP/TIFF) "
       "لأن إعادة ضغط JPEG تدمر بتات LSB. الترويسة تُكتب في مواضع ثابتة والجسم يُوزع بترتيب مشتق من "
       "SHA-256(ملح + المفتاح) مع تقنيع XOR — مفتاح خاطئ يعطي خطأ صريحًا.", bullet=True, size=10.5)
IMG(doc, "s11.png", "الإخفاء في PNG بمحرك LSB الأصلي")
REAL(doc, "s11.png")
IMG(doc, "s12.png", "الاستخراج — تطابق بايت-ببايت مع الأصل")
REAL(doc, "s12.png")
H(doc, "التقنية 2: Steghide (أداة خارجية INTEGRATED)", size=12)
P(doc, "أداة حقيقية تعمل في المجال التحويلي بخوارزمية نظرية المخططات مع ضغط وتشفير داخلي بكلمة مرور؛ "
       "تدعم JPEG/BMP/WAV ولا تدعم PNG (فشل الصيغة يظهر بصدق).", bullet=True, size=10.5)
IMG(doc, "s13.png", "الإخفاء في JPEG عبر steghide الحقيقية")
REAL(doc, "s13.png")

# ---- 6.2 audio
H(doc, "6.2 الإخفاء في الصوت — ثلاث تقنيات", size=14, color=CYAN)
H(doc, "التقنية 1: LSB على عينات PCM", size=12)
P(doc, "استبدال البت الأقل أهمية في عينات 16-بت (الخطوة = 1/32768 من المدى — دون عتبة السمع). "
       "الصيغ الأخرى (MP3/FLAC) تُحول أولًا إلى WAV PCM عبر FFmpeg.", bullet=True, size=10.5)
IMG(doc, "s15.png", "الإخفاء في الصوت LSB — SNR مقاس")
REAL(doc, "s15.png")
H(doc, "التقنية 2: ترميز الطور (Phase Coding)", size=12)
P(doc, "تقسيم الإشارة إلى مقاطع 1024 عينة ← FFT ← فرض فرق طور +π/4 (بت 0) أو −π/4 (بت 1) على حزمة عالية "
       "الطاقة ← إعادة البناء بنفس المقادير (فلا يلاحظ السمع). يحتاج حاملًا ≥ 15 ثانية وأقوى ضد التحليل "
       "البسيط، لكنه هش لإعادة الترميز.", bullet=True, size=10.5)
IMG(doc, "s17.png", "الإخفاء بترميز الطور على حامل 16 ثانية")
REAL(doc, "s17.png")
H(doc, "التقنية 3: طيف الانتشار DSSS", size=12)
P(doc, "كل بت ينتشر عبر 1024 خانة بمتتالية ضوضاء ±1 مشتقة من المفتاح (كسب معالجة مقاس بالـ dB)، "
       "والاستخراج **أعمى**: بالمفتاح فقط دون الحامل الأصلي — طاقة الرسالة تحت أرضية الضوضاء.", bullet=True, size=10.5)
IMG(doc, "s19.png", "الإخفاء بـ DSSS على حامل 30 ثانية")
REAL(doc, "s19.png")
IMG(doc, "s20.png", "الاستخراج الأعمى (بالمفتاح فقط) — تطابق كامل")
REAL(doc, "s20.png")

# ---- 6.3 video
H(doc, "6.3 الإخفاء في الفيديو — تقنيتان", size=14, color=CYAN)
H(doc, "التقنية 1: LSB على إطارات FFV1 (خط بلا فقد)", size=12)
P(doc, "ffprobe ← تفكيك الإطارات إلى PNG عبر ffmpeg ← تضمين LSB إطارًا بإطار (ترويسة في الإطار 0 وجسم "
       "بترتيب مفتاحي) ← إعادة ترميز FFV1 داخل MKV (فائد صفر) ← تحقق تلقائي بإعادة الاستخراج. "
       "الأداة ترفض إخراج mp4 مقصودًا: أي ضغط فائد يدمر LSB.", bullet=True, size=10.5)
IMG(doc, "s21.png", "الإخفاء في فيديو عبر FFV1 LSB")
REAL(doc, "s21.png")
H(doc, "التقنية 2: حاوية EOF المخصصة (مشفرة)", size=12)
P(doc, "بنية: بصمة SNXEOF0001 + إصدار + ملح 16B + nonce 12B + اسم الملف + نص مشفر **AES-256-GCM** "
       "(مشتق المفتاح PBKDF2-SHA256 بـ 200,000 تكرار)، وتُلحق بعد نهاية الفيديو مع trailer يحمل الطول. "
       "المشغل يقرأ الفيديو طبيعيًا 100%، وDetect يكشفها من آخر 18 بايت دون مفتاح. تصريح صريح: "
       "حاوية أكاديمية خاصة — **ليست OpenPuff ولا متوافقة معها**.", bullet=True, size=10.5)
IMG(doc, "s23.png", "الإخفاء بحاوية EOF بعد نهاية الفيديو")
REAL(doc, "s23.png")
IMG(doc, "s24.png", "كشف الحاوية (Detect) — من ذيل الملف دون مفتاح")
REAL(doc, "s24.png")

# ---- 6.4 network
H(doc, "6.4 الإخفاء في الشبكة — قناة IPv4 Identification", size=14, color=CYAN)
H(doc, "المبدأ والتقنية", size=12)
P(doc, "حقل Identification في ترويسة IPv4 تحدده المصدر ولا توحّده الراوترات ← يمكن حمل **15 بت في كل حزمة** "
       "دون أن يتغير في العبور. تُؤطَّر الحمولة (بصمة SNXNET1 + عدد القطع + طول + CRC-32) وتُقسم إلى قطع "
       "15-بت، والبت 15 يحمل parity على فهرس القطعة لاكتشاف الفقد وإعادة الترتيب عند إعادة التجميع.", bullet=True, size=10.5)
P(doc, "ثلاث أدوات في الشاشة: **Preview** (خريطة قيم IP-ID دون إرسال) · **Send/Receive** (قناة UDP أساسية غير "
       "سرية — للمقارنة التعليمية) · **القناة السرية IPv4-ID** (تتطلب root + وضع المختبر المصرَّح — أثبتت PASS "
       "في preflight). والوجه الجنائي: **مفسر PCAP** يعلن أن «الشذوذ مؤشر لا دليل».", bullet=True, size=10.5)
IMG(doc, "s26.png", "خريطة حقن حقل IPv4-ID لحمولة حقيقية")
REAL(doc, "s26.png")
IMG(doc, "s28.png", "إرسال واستقبال فعلي على loopback — تطابق الحمولة")
REAL(doc, "s28.png")

# ---- 6.5 malware
H(doc, "6.5 البرامج الضارة والفيروسات — التحليل الساكن + الإخفاء التنفيذي (تقنيتان)", size=14, color=CYAN)
H(doc, "أ) التحليل الساكن الدفاعي (لا تنفيذ إطلاقًا)", size=12)
P(doc, "تجزئات ← نوع الملف ← سلاسل ← إنتروبيا ← فحص PE كامل (أقسام الحزم المعروفة UPX/Themida، إنتروبيا قسم "
       "≥ 7.2، نقطة الدخول في آخر قسم تنفيذي، توقيت ترجمة مستحيل) ← IOC ← انتحال الامتدادات وRTLO ← "
       "timestomping ← التمويه ← PyInstaller ← الأرشيفات ← ELF ← Office. الدرجة = 8×عدد المؤشرات (سقف 100) "
       "والثقة بالارتباط: مؤشر واحد = Low · أقل من 3 = Medium · 3+ = High — «مؤشر واحد لا يثبت شيئًا».", bullet=True, size=10.5)
IMG(doc, "s29.png", "تحليل ساكن كامل داخل الواجهة — دون تنفيذ العينة")
REAL(doc, "s29.png")
H(doc, "ب) الإخفاء التنفيذي — التقنية 1: PE Overlay", size=12)
P(doc, "إلحاق حاوية AES-256-GCM (بصمة SNXPEST001، PBKDF2 200k) بعد نهاية آخر قسم في PE + trailer بالطول — "
       "نمط dropper/installer الكلاسيكي. الأثر الجنائي: الحجم يزيد ← يكشفها مقارنة نهاية جدول الأقسام بالحجم.", bullet=True, size=10.5)
IMG(doc, "s34.png", "تقنية Overlay: الزيادة في الحجم ظاهرة ومكتشفة")
REAL(doc, "s34.png")
H(doc, "ج) الإخفاء التنفيذي — التقنية 2: Section Slack", size=12)
P(doc, "كتابة الحاوية داخل حشو القسم (الفرق بين VirtualSize وSizeOfRawData) وملء الباقي أصفارًا — "
       "**الحجم لا يتغير إطلاقًا** فيكشفها فقط تحليل جدول الأقسام وفحص الحشو بايت-ببايت (زر Scan hidden data)، "
       "وإعادة التحليل الساكن تضع مؤشر Executable stego تلقائيًا.", bullet=True, size=10.5)
IMG(doc, "s30.png", "إخفاء 73 بايت في حشو القسم — الحجم لم يتغير (2048←2048)")
REAL(doc, "s30.png")
IMG(doc, "s31.png", "الكشف الساكن: تحديد القسم .text وبداية الغطس وحجمه")
REAL(doc, "s31.png")
IMG(doc, "s32.png", "الاستخراج — تطابق SHA-256 مع الأصل")
REAL(doc, "s32.png")
IMG(doc, "s33.png", "التحليل الساكن يلوّح بالغطس تلقائيًا — حلقة مكتملة")
REAL(doc, "s33.png")
P(doc, "حامل التدريب: PE صناعي بترويسات صحيحة وقسم .text بحشو 768 بايت وبلا أي كود (نقطة دخول 0) — "
       "نعلّم التقنية على عينة لا يمكنها العمل أصلًا (أمانة منهجية).", size=10.5, color=AMBER)

# ---- 6.0 text (bonus)
H(doc, "6.0 الإخفاء في النص — LSB مع مفتاح (إضافة)", size=14, color=CYAN)
P(doc, "أحرف مؤهلة فقط (قلب LSB يبقى حرفًا/رقمًا) ← إطار (بصمة SNX1 + فحص مفتاح + طول + CRC-32) ← ضغط zlib "
       "← تقنيع XOR بمجرى SHA-256-CTR ← كتابة بترتيب Fisher-Yates محكوم بالمفتاح. الكشف: الانحياز LSB "
       "بزر Analyse.", bullet=True, size=10.5)
IMG(doc, "s09.png", "إخفاء رسالة حقيقية داخل نص الغلاف")
REAL(doc, "s09.png")

# ============ 7 security
doc.add_page_break()
H(doc, "7. الأمان والأخلاقيات", size=17)
TBL(doc, ["المبدأ", "التطبيق في الأداة"],
    [["قراءة-فقط افتراضيًا", "الدليل الأصلي لا يُعدَّل أبدًا؛ كل العمليات على نسخ؛ إثبات قبل/بعد بـ SHA-256"],
     ["برمجيات خبيثة: ساكن فقط", "العينة لا تُنفَّذ إطلاقًا؛ لا sandbox ولا أي كود هجومي في المشروع"],
     ["شبكة مقيدة", "الإرسال مقفول افتراضيًا (مفتاح مصرَّح + root للقناة السرية) والوجهات loopback/RFC1918"],
     ["لا نتائج ملفقة", "أداة غائبة = UNAVAILABLE · الشذوذ مؤشر لا دليل · العينات موسومة TRAINING-SYNTHETIC"],
     ["تنفيذ آمن للأدوات", "ToolExecutor بقائمة وسائط (بدون shell=True) ومهلة لكل أداة — يمنع حقن الأوامر"]],
    widths=[1.9, 4.9])

# ============ 8 verification
H(doc, "8. الاختبارات والتحقق من الصحة", size=17)
TBL(doc, ["الطقم", "النتيجة"],
    [["pytest (اختبارات الانحدار)", "34 ناجحة — تشمل جولات إخفاء/استخراج لكل تقنية واختبار القضايا A–D"],
     ["preflight_check.py", "VERDICT: READY (21 PASS / 1 GATED / 1 BLOCKED بتصميم) — وبصلاحيات root: 22 PASS مع تأكيد قناة IPv4-ID"],
     ["gui_smoke_test", "20/20 واجهة تُرسم بلا أخطاء + لقطة إثبات للإخفاء التنفيذي"],
     ["gui_button_test", "11/11 مسار أزرار حقيقي (منها Hide/Scan/Extract التنفيذية)"],
     ["demo_end_to_end", "كل خطوات العرض نُفذت فعليًا وتولد تقارير حقيقية"],
     ["التشغيل الموثق", "36 خطوة بقياساتها ولقطاتها (هذه الوثيقة مأخوذة منها) — كل جولات الإخفاء 10/10 بتطابق SHA-256"]],
    widths=[2.3, 4.5])

# ============ 9 mapping
H(doc, "9. خريطة المتطلبات ← التنفيذ ← التحقق", size=17)
TBL(doc, ["متطلب المقرر", "الشاشة", "التقنية/الأداة", "التحقق"],
    [["عرض البيانات الوصفية", "Metadata ← Read", "ExifTool", "32 وسمًا مقروءًا فعليًا (s05)"],
     ["حقن البيانات الوصفية", "Metadata ← Inject/Strip/Compare", "ExifTool على نسخة", "3 وسوم + unchanged=True (s06/s07)"],
     ["القضية 1 (صورة)", "Cases + Image", "LSB أصلي + Steghide", "استرجاع مطابق SHA-256 (s11–s14)"],
     ["القضية 2 (صوت)", "Audio", "LSB + Phase + DSSS", "SNR ~97dB + تطابق (s15–s20)"],
     ["القضية 3 (شبكة/خبيث)", "Network + Malware", "IPv4-ID + PCAP + Static", "SUCCESS/SUCCESS + High (s26–s29)"],
     ["إخفاء الفيديو", "Video", "FFV1 LSB + EOF container", "تطابق + كشف (s21–s25)"],
     ["البرامج الضارة/الفيروسات", "Malware", "Overlay + Slack + كشف ساكن", "حجم ثابت + تطابق + إلواح (s29–s34)"],
     ["التوثيق", "Reports + Logs", "PDF/HTML/JSON + سجل تدقيق", "3 صيغ من السجلات (s35/s36)"]],
    widths=[1.7, 1.7, 1.9, 1.5], size=9.5)

# ============ 10 install
H(doc, "10. التثبيت والتشغيل وإعادة إنتاج النتائج", size=17)
P(doc, "على Kali/Debian:", size=11, bold=True)
P(doc, "sudo ./scripts/install_kali.sh    # الأدوات الخارجية + Qt + .venv + الاعتماديات", size=10, mono=True)
P(doc, "./scripts/preflight.sh             # فحص شامل: VERDICT: READY", size=10, mono=True)
P(doc, "./scripts/run.sh                   # الواجهة الرسومية", size=10, mono=True)
P(doc, "sudo ./scripts/run.sh              # مع قناة IPv4-ID السرية (raw sockets)", size=10, mono=True)
P(doc, "إعادة إنتاج التشغيل الموثق في هذه الوثيقة:", size=11, bold=True)
P(doc, "QT_QPA_PLATFORM=offscreen python scripts/capture_evidence_report.py", size=10, mono=True)
P(doc, "python scripts/build_evidence_report_ar.py     # التقرير المصوَّر HTML", size=10, mono=True)
P(doc, "python -m pytest -q                            # 34 اختبارًا", size=10, mono=True)

# ============ 11 conclusion
H(doc, "11. الخاتمة", size=17)
P(doc, "حقق StegoNexus تغطية كاملة لمتطلبات المقرر: أداة عرض وحقن البيانات الوصفية، وأدوات الإخفاء والاستخراج "
       "العاملة داخل قضايا كاملة الأركان، وأدوات إخفاء واستخراج لكل الوسائط (الصورة تقنيتان، الصوت ثلاث، "
       "الفيديو تقنيتان، الشبكة قناة سرية مع مفسر، والبرامج الضارة والفيروسات بتقنيتي Overlay وSlack مع كشف "
       "ساكن مكتمل) — وكل ذلك موثق بقياسات حقيقية قابلة لإعادة التشغيل بأمر واحد. القيمة الأكاديمية الجوهرية: "
       "**منصة تقول الحقيقة دائمًا عن ما فعلته وبأي أداة وبأي تجزئة — وما لا تستطيعه تعلنه بوضوح.**", size=11.5)
doc.add_paragraph()
P(doc, "تطوير وإعداد: Mohammed Moneer Al-absi · إشراف: أسامة الشلالي · 2026", size=11, color=GRAY)

doc.save(OUT)
print("saved", OUT)
